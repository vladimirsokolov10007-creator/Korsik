"""Convert approved source files to UTF-8 Markdown for Korsik's local RAG index.

Usage from the project root:
    python backend/scripts/ingest_documents.py "C:\\path\\to\\approved_docs" backend/data/knowledge_base

For scanned PDFs this script intentionally stops with a warning: OCR must be reviewed by
an owner before material becomes a source of truth for an HR assistant.
"""
from __future__ import annotations

import re
import shutil
import sys
from pathlib import Path


def safe_name(path: Path) -> str:
    return re.sub(r"[^0-9A-Za-zА-Яа-яЁё_-]+", "_", path.stem).strip("_") or "source"


def extract_pdf(path: Path) -> str:
    try:
        import fitz  # PyMuPDF
    except ImportError as error:
        raise RuntimeError("Для PDF установите зависимости из requirements.txt") from error
    document = fitz.open(path)
    try:
        pages = [page.get_text("text").strip() for page in document]
    finally:
        document.close()
    text = "\n\n".join(f"## Страница {index}\n\n{page}" for index, page in enumerate(pages, start=1) if page)
    if len(text) < 100:
        raise RuntimeError("PDF выглядит сканированным или пустым. Сделайте OCR и проверьте текст перед загрузкой в RAG.")
    return text


def extract_xlsx(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError("Для XLSX установите зависимости из requirements.txt") from error
    workbook = load_workbook(path, read_only=True, data_only=True)
    parts: list[str] = []
    try:
        for worksheet in workbook.worksheets:
            rows = []
            for row in worksheet.iter_rows(values_only=True):
                values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                if values:
                    rows.append(" | ".join(values))
                if len(rows) >= 400:
                    rows.append("[Ограничение MVP: показаны первые 400 непустых строк листа]")
                    break
            if rows:
                parts.append(f"## Лист: {worksheet.title}\n\n" + "\n".join(rows))
    finally:
        workbook.close()
    return "\n\n".join(parts)


def convert(source: Path) -> str:
    suffix = source.suffix.lower()
    if suffix in {".md", ".txt"}:
        return source.read_text(encoding="utf-8", errors="replace")
    if suffix == ".pdf":
        return extract_pdf(source)
    if suffix in {".xlsx", ".xlsm"}:
        return extract_xlsx(source)
    raise RuntimeError(f"Неподдерживаемый формат: {source.suffix}")


def main() -> int:
    if len(sys.argv) != 3:
        print(__doc__.strip())
        return 2
    source_dir = Path(sys.argv[1]).expanduser().resolve()
    output_dir = Path(sys.argv[2]).expanduser().resolve()
    if not source_dir.is_dir():
        print(f"Не найдена папка источников: {source_dir}")
        return 2
    output_dir.mkdir(parents=True, exist_ok=True)
    allowed = {".md", ".txt", ".pdf", ".xlsx", ".xlsm"}
    created = 0
    for source in sorted(path for path in source_dir.iterdir() if path.is_file() and path.suffix.lower() in allowed):
        try:
            body = convert(source)
            header = f"# {source.stem}\n\n> Импортировано из утверждённого источника: {source.name}. Проверьте текст и актуальность перед запуском.\n\n"
            target = output_dir / f"{safe_name(source)}.md"
            target.write_text(header + body, encoding="utf-8")
            print(f"OK: {source.name} → {target.name}")
            created += 1
        except RuntimeError as error:
            print(f"SKIP: {source.name}: {error}")
    print(f"Готово: {created} файл(а/ов). Затем вызовите POST /api/admin/reload-knowledge или перезапустите сервис.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
